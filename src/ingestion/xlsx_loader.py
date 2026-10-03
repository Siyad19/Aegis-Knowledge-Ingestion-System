from openpyxl import load_workbook


def load_xlsx(file_path):

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    content = []

    for sheet in workbook.worksheets:

        for row in sheet.iter_rows():

            values = []

            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))

            if values:

                content.append({
                    "location": {
                        "sheet": sheet.title,
                        "row": row[0].row
                    },
                    "text": " | ".join(values),
                    "extraction_method": "spreadsheet"
                })

    return content