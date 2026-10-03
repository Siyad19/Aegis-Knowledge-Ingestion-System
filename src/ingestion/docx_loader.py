from docx import Document


def load_docx(file_path):

    document = Document(file_path)

    content = []

    for index, paragraph in enumerate(document.paragraphs, start=1):

        text = paragraph.text.strip()

        if text:
            content.append({
                "location": {
                    "paragraph": index
                },
                "text": text,
                "extraction_method": "text"
            })

    # Extract tables
    for table_index, table in enumerate(document.tables, start=1):

        for row_index, row in enumerate(table.rows, start=1):

            row_data = []

            for cell in row.cells:
                row_data.append(cell.text.strip())

            content.append({
                "location": {
                    "table": table_index,
                    "row": row_index
                },
                "text": " | ".join(row_data),
                "extraction_method": "table"
            })

    return content