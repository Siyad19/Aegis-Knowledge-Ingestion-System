from pptx import Presentation


def load_pptx(file_path):

    presentation = Presentation(file_path)

    content = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        for shape_number, shape in enumerate(
            slide.shapes,
            start=1
        ):

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:

                    content.append({
                        "location": {
                            "slide": slide_number,
                            "shape": shape_number
                        },
                        "text": text,
                        "extraction_method": "pptx"
                    })

    return content